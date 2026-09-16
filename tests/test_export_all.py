"""The one-click "export everything" workbook. Uses fake data, no network."""

import io

import pytest
from openpyxl import load_workbook

from backend import exports
from core import positions, realized


def test_workbook_writes_and_reads_back():
    sheets = {
        "Holdings": [["symbol", "shares", "cost_basis"], ["NVDA", 400, 178.5]],
        "About": [["section", "status"], ["Holdings", "1 lot"]],
    }
    wb = load_workbook(io.BytesIO(exports.sheets_to_xlsx(sheets)))
    assert wb.sheetnames == ["Holdings", "About"]
    assert [c.value for c in wb["Holdings"][2]] == ["NVDA", 400, 178.5]


def test_header_row_is_frozen_and_bold():
    wb = load_workbook(io.BytesIO(exports.sheets_to_xlsx(
        {"S": [["a", "b"], [1, 2]]})))
    ws = wb["S"]
    assert ws.freeze_panes == "A2"
    assert ws["A1"].font.bold


def test_sheet_names_are_sanitised():
    r"""Excel rejects []:*?/\ and caps names at 31 characters."""
    long_name = "Performance " + "x" * 40
    wb = load_workbook(io.BytesIO(exports.sheets_to_xlsx(
        {long_name: [["a"], [1]], "bad[]:*?/name": [["a"], [1]]})))
    for name in wb.sheetnames:
        assert len(name) <= 31
        assert not any(ch in name for ch in "[]:*?/\\")


def test_percent_and_money_columns_get_number_formats():
    wb = load_workbook(io.BytesIO(exports.sheets_to_xlsx(
        {"S": [["pnl_pct", "market_value", "symbol"], [10.5, 1234.5, "X"]]})))
    ws = wb["S"]
    assert "%" in ws["A2"].number_format
    assert "#,##0.00" in ws["B2"].number_format
    assert ws["C2"].number_format in ("General", "@")


def test_empty_sheet_does_not_crash():
    wb = load_workbook(io.BytesIO(exports.sheets_to_xlsx({"Empty": []})))
    assert "Empty" in wb.sheetnames


def test_realized_rows_round_trip_into_a_sheet():
    sold = realized.parse_sales([
        {"symbol": "TSLA", "shares": 50, "cost_basis": 190.0, "exit_price": 180.0,
         "opened": "2026-09-16 09:45", "closed": "2026-09-16 11:20"}])
    wb = load_workbook(io.BytesIO(exports.sheets_to_xlsx(
        {"Realized Gains": realized.to_csv_rows(sold)})))
    header = [c.value for c in wb["Realized Gains"][1]]
    assert "holding_label" in header and "term" in header
    row = [c.value for c in wb["Realized Gains"][2]]
    # The clock time on an intraday trade must survive the round trip.
    assert "09:45" in str(row[header.index("opened")])


def test_dicts_to_rows_restricts_and_orders_columns():
    rows = exports.dicts_to_rows(
        [{"symbol": "A", "pnl": 1, "extra": "ignored"}], ["symbol", "pnl"])
    assert rows[0] == ["symbol", "pnl"]
    assert rows[1] == ["A", 1]


def test_missing_values_become_blank_not_none():
    rows = exports.dicts_to_rows([{"symbol": "A"}], ["symbol", "pnl"])
    assert rows[1] == ["A", ""]


def test_positions_csv_rows_have_the_documented_columns():
    parsed = positions.parse_positions(
        [{"symbol": "NVDA", "shares": 400, "cost_basis": 178.5}])
    rows = positions.to_csv_rows(
        positions.value_portfolio(parsed, [])["positions"])
    assert rows[0] == positions.CSV_COLUMNS
    assert rows[1][0] == "NVDA"
